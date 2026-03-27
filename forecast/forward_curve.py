"""
Forward curve loading.

Mirrors VBA worksheet module: forward_curve.cls

Loads the forward SOFR/LIBOR curve from an Intex Cashflows report,
used to determine interest rate assumptions for scenario analysis.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd


def load_from_workbook(wb_path: str | Path) -> pd.DataFrame:
    """Load forward curve from the Forecast workbook's 'Forward Curve' sheet.

    Returns:
        DataFrame with columns: Period, Forward 3-Month LIBOR
    """
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    ws = wb["Forward Curve"]

    # Data starts at row 5: Period (A), Forward 3-Month LIBOR (B)
    rows = []
    row_num = 5
    while True:
        period = ws.cell(row=row_num, column=1).value
        if period is None:
            break
        rate = ws.cell(row=row_num, column=2).value
        rows.append({"Period": int(period), "Forward 3-Month LIBOR": float(rate)})
        row_num += 1

    wb.close()
    df = pd.DataFrame(rows)
    print(f"  Forward curve loaded: {len(df)} periods.")
    return df


def load_from_cashflows_report(cashflows_path: str | Path) -> pd.DataFrame:
    """Extract forward curve from the rate vector string in an Intex Cashflows report.

    Mirrors VBA method: forward_curve.Load()

    The first non-"Total" worksheet contains a row starting with "LIBOR (3mo)"
    where column D has the full rate vector as a space-delimited string.
    """
    wb = openpyxl.load_workbook(cashflows_path, read_only=True, data_only=True)

    # Find the first suitable worksheet.
    ws = None
    for sheet in wb.worksheets:
        if not sheet.title.startswith("Total"):
            ws = sheet
            break
    if ws is None:
        ws = wb.worksheets[0]

    # Find the "LIBOR (3mo)" row.
    row_num = 1
    while row_num < 500:
        if ws.cell(row=row_num, column=1).value == "LIBOR (3mo)":
            break
        row_num += 1

    rate_vector_str = str(ws.cell(row=row_num, column=4).value or "")
    wb.close()

    if not rate_vector_str:
        print("  WARNING: Could not find forward curve in cashflows report.")
        return pd.DataFrame(columns=["Period", "Forward 3-Month LIBOR"])

    rates = rate_vector_str.split()
    rows = [
        {"Period": i + 1, "Forward 3-Month LIBOR": float(r)}
        for i, r in enumerate(rates)
    ]

    df = pd.DataFrame(rows)
    print(f"  Forward curve extracted: {len(df)} periods.")
    return df


def is_upward_sloping(curve_df: pd.DataFrame, start_period: int, periods: int) -> bool:
    """Check if the curve is upward sloping over the specified window.

    Mirrors VBA method: forward_curve.UpwardSloping()
    """
    end_period = start_period + periods
    curve = curve_df.set_index("Period")["Forward 3-Month LIBOR"]

    start_rate = curve.get(start_period, 0)
    end_rate = curve.get(end_period, 0)

    return end_rate > start_rate
