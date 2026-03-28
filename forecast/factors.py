"""
Forecasted factor call date extraction.

Mirrors VBA worksheet module: forecasted_factors.cls

Step 3: Clear forecasted factors (trivial — just start with empty data).
Step 5: Extract factor call dates from preliminary Intex Cashflows report.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Optional

import openpyxl

from forecast.config import CONFIG
from forecast.models import Tranche


class ForecastedFactors:
    """Lookup table for factor-based call dates, keyed by (scenario_name, deal_name).

    When the deal's collateral balance drops below 35% of the original deal balance,
    the deal is expected to be redeemed. This class extracts those dates from a
    preliminary Intex Cashflows report.

    Mirrors VBA class: forecasted_factors.cls
    """

    def __init__(self):
        self._data: dict[tuple[str, str], date] = {}

    def clear(self) -> None:
        """Clear all factor call dates (Step 3)."""
        self._data.clear()

    def get_factor_call_date(self, scenario_name: str, deal_name: str) -> Optional[date]:
        """Look up the factor call date for a scenario/deal combination.

        Returns None if no factor call date exists (equivalent to VBA returning 'N/A').
        """
        return self._data.get((scenario_name, deal_name))

    def add(self, scenario_name: str, deal_name: str, call_date: date) -> None:
        """Add a factor call date entry."""
        self._data[(scenario_name, deal_name)] = call_date

    @property
    def count(self) -> int:
        return len(self._data)

    def import_from_cashflows_report(
        self,
        cashflows_path: str | Path,
        tranches: list[Tranche],
        factor_threshold: float | None = None,
    ) -> None:
        """Extract factor call dates from a preliminary Intex Cashflows Excel export.

        For each COLLAT worksheet, finds the first period where the deal balance
        drops below ``factor_threshold`` of the original deal balance.

        Mirrors VBA method: forecasted_factors.ImportFactorCallDates()

        Args:
            cashflows_path: Path to the Intex Cashflows Excel export.
            tranches: List of Tranche objects (used for original deal balances).
            factor_threshold: Factor below which a deal is considered called.
                              Defaults to CONFIG.factor_call.factor_threshold.
        """
        if factor_threshold is None:
            factor_threshold = CONFIG.factor_call.factor_threshold
        self.clear()

        # Build lookup of original deal balances from tranches.
        original_balances: dict[str, float] = {}
        for t in tranches:
            if t.intex_deal_name and t.intex_deal_name not in original_balances:
                original_balances[t.intex_deal_name] = t.orig_deal_balance

        print(f"Importing factor call dates from '{Path(cashflows_path).name}'...")
        wb = openpyxl.load_workbook(cashflows_path, read_only=True, data_only=True)

        for ws in wb.worksheets:
            ws_name = ws.title

            # Only process COLLAT worksheets.
            if "COLLAT" not in ws_name:
                continue

            # Read deal name from worksheet.
            deal_name = ""
            row_num = 1
            while row_num <= 500:
                if ws.cell(row=row_num, column=1).value == "Deal Name:":
                    deal_name = str(ws.cell(row=row_num, column=2).value or "")
                    break
                row_num += 1

            if not deal_name:
                print(f"  WARNING: Could not find 'Deal Name:' in '{ws_name}', skipping.")
                continue

            orig_balance = original_balances.get(deal_name, 0)
            if orig_balance <= 0:
                print(f"  WARNING: No original deal balance for '{deal_name}', skipping.")
                continue

            # Find header row ("Period") and "Deal Balance" column.
            while row_num < 500:
                if ws.cell(row=row_num, column=1).value == "Period":
                    break
                row_num += 1

            balance_col = 1
            while balance_col < 200:
                if ws.cell(row=row_num, column=balance_col).value == "Deal Balance":
                    break
                balance_col += 1

            # Advance to first data row (3 rows below header).
            data_row = row_num + 3
            factor_call_date: Optional[date] = None

            while ws.cell(row=data_row, column=2).value is not None:
                balance_val = ws.cell(row=data_row, column=balance_col).value
                if balance_val is not None and float(balance_val) > 0:
                    if float(balance_val) / orig_balance < factor_threshold:
                        date_val = ws.cell(row=data_row, column=2).value
                        factor_call_date = _to_date(date_val)
                        break
                data_row += 1

            # Read scenario name from "User Comment 2".
            scenario_name = ""
            while row_num < 1000:
                if ws.cell(row=row_num, column=1).value == "User Comment 2":
                    scenario_name = str(ws.cell(row=row_num, column=4).value or "")
                    break
                row_num += 1

            # Fall back to deducing scenario name from worksheet name.
            if not scenario_name:
                scenario_name = _get_scenario_name_from_ws_name(ws_name)

            if factor_call_date is not None and scenario_name:
                self.add(scenario_name, deal_name, factor_call_date)

        wb.close()
        print(f"  {self.count} factor call dates extracted.")

    def load_from_workbook(self, wb_path: str | Path) -> None:
        """Load pre-existing factor call dates from the Forecast workbook's 'Forecasted Factors' sheet."""
        self.clear()

        wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
        ws = wb["Forecasted Factors"]

        # Data starts at row 7: Scenario Name (A), Deal Name (B), Factor Call Date (C)
        row_num = 7
        while True:
            scenario = ws.cell(row=row_num, column=1).value
            if scenario is None:
                break
            deal = str(ws.cell(row=row_num, column=2).value or "")
            call_date_val = ws.cell(row=row_num, column=3).value
            if deal and call_date_val:
                self.add(str(scenario), deal, _to_date(call_date_val))
            row_num += 1

        wb.close()
        print(f"  {self.count} factor call dates loaded from workbook.")

    def to_dataframe(self):
        """Export factor call dates as a DataFrame."""
        import pandas as pd

        rows = []
        for (scenario, deal), call_date in sorted(self._data.items()):
            rows.append({
                "Scenario Name": scenario,
                "Deal Name": deal,
                "Factor Call Date": call_date,
            })
        return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Helper functions
# ---------------------------------------------------------------------------

def _to_date(val) -> date:
    """Convert a value to a date object."""
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, (int, float)):
        # Intex sometimes outputs dates as YYYYMMDD integers.
        s = str(int(val))
        return date(int(s[:4]), int(s[4:6]), int(s[6:8]))
    if isinstance(val, str):
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y%m%d"):
            try:
                return datetime.strptime(val, fmt).date()
            except ValueError:
                continue
    raise ValueError(f"Cannot convert {val!r} to date")


def _get_scenario_name_from_ws_name(ws_name: str) -> str:
    """Deduce scenario name from Intex Cashflows worksheet name.

    Handles: "DEAL-COLLAT-75 bps Tightening", "DEAL-COLLAT-Flat", etc.
    """
    last_hyphen = ws_name.rfind("-")
    if last_hyphen < 0:
        return ws_name
    raw = ws_name[last_hyphen + 1:].strip()

    if "Flat" in raw:
        return "Flat"
    if raw.startswith("CCC"):
        return "CCC_Stress"
    if raw.startswith("High"):
        return "High_Prepay"

    parts = raw.split()
    if len(parts) >= 3:
        num_bps = f"{parts[0]} {parts[1]}"
        first_letter = parts[2][0].upper()
        if first_letter == "T":
            return f"{num_bps} Tightening"
        elif first_letter == "W":
            return f"{num_bps} Widening"

    return raw
