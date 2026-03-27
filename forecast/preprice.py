"""
Pre-price deals lookup.

Mirrors VBA worksheet module: preprice_deals.cls

Pre-price deals are CLO tranches that haven't priced yet and require special
Intex identifiers (deal name + password) for cashflow modeling.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd


class PrePriceDeals:
    """Lookup table for pre-price deal information, keyed by CUSIP.

    Mirrors the VBA functions: GetIdentifier, GetName, GetDealName, GetPassword
    """

    def __init__(self, data: dict[str, dict] | None = None):
        self._data: dict[str, dict] = data or {}

    @classmethod
    def from_forecast_workbook(cls, wb_path: str | Path) -> PrePriceDeals:
        """Load pre-price deals from the Forecast workbook's 'Pre-Price Deals' sheet."""
        import openpyxl

        wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
        ws = wb["Pre-Price Deals"]

        # Headers at row 3: CUSIP, Pre-Price Deal Name, Pre-Price Tranche Name, Password, Identifier String
        data: dict[str, dict] = {}
        row_num = 4
        while True:
            cusip = ws.cell(row=row_num, column=1).value
            if cusip is None:
                break
            cusip = str(cusip).strip()
            data[cusip] = {
                "deal_name": str(ws.cell(row=row_num, column=2).value or ""),
                "tranche_name": str(ws.cell(row=row_num, column=3).value or ""),
                "password": str(ws.cell(row=row_num, column=4).value or ""),
                "identifier": str(ws.cell(row=row_num, column=5).value or ""),
            }
            row_num += 1

        wb.close()
        print(f"  {len(data)} pre-price deals loaded.")
        return cls(data)

    @classmethod
    def from_dataframe(cls, df: pd.DataFrame) -> PrePriceDeals:
        """Load pre-price deals from a DataFrame with columns:
        CUSIP, Pre-Price Deal Name, Pre-Price Tranche Name, Password, Identifier String
        """
        data: dict[str, dict] = {}
        for _, row in df.iterrows():
            cusip = str(row.get("CUSIP", "")).strip()
            if not cusip:
                continue
            data[cusip] = {
                "deal_name": str(row.get("Pre-Price Deal Name", "") or ""),
                "tranche_name": str(row.get("Pre-Price Tranche Name", "") or ""),
                "password": str(row.get("Password", "") or ""),
                "identifier": str(row.get("Identifier String", "") or ""),
            }
        return cls(data)

    def get_identifier(self, cusip: str) -> str:
        """Return the Intex identifier string for the CUSIP, or 'N/A'."""
        entry = self._data.get(cusip)
        if entry and entry.get("identifier"):
            return entry["identifier"]
        return "N/A"

    def get_name(self, cusip: str) -> str:
        """Return the Intex name (DEAL,TRANCHE) for the CUSIP, or 'N/A'."""
        entry = self._data.get(cusip)
        if entry and entry.get("deal_name"):
            return f"{entry['deal_name']},{entry['tranche_name']}".upper()
        return "N/A"

    def get_deal_name(self, cusip: str) -> str:
        """Return the Intex deal name for the CUSIP, or 'N/A'."""
        entry = self._data.get(cusip)
        if entry and entry.get("deal_name"):
            return entry["deal_name"].upper()
        return "N/A"

    def get_password(self, cusip: str) -> str:
        """Return the pre-price password for the CUSIP, or 'N/A'."""
        entry = self._data.get(cusip)
        if entry and entry.get("password"):
            return str(entry["password"]).upper()
        return "N/A"
